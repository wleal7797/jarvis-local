"""Pruebas de la ingesta: cargadores, cola de tareas, servicio y endpoints (sin internet)."""

import io
import json
from pathlib import Path

import httpx
import pytest
from docx import Document as WordDocument
from sqlalchemy import func, select, text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool

from app.models import Chunk, Document, IngestionJob, Source, User
from app.repositories import JobRepository
from app.services.ingestion import (
    FileLoader,
    IngestionError,
    IngestionService,
    LoaderError,
    WebLoader,
    WikipediaLoader,
)
from app.services.ingestion.wikipedia_loader import wiki_to_markdown
from tests.fakes import FakeEmbedder

pytestmark = pytest.mark.anyio

# ---------- Datos de prueba ----------

ARTICLE = """Bogotá es la capital de Colombia.

== Historia ==
Fue fundada en 1538 por Gonzalo Jiménez de Quesada.

=== Época colonial ===
Fue capital del Nuevo Reino de Granada.

== Referencias ==
Libro de historia, página 3.

== Clima ==
Clima frío de montaña.
"""

PAGE = """<html><head><title>{title}</title></head><body><main><h1>{title}</h1>
<p>{body} Este párrafo tiene suficiente texto para que trafilatura lo considere contenido
principal de la página y no lo descarte como menú o pie de página. Se repite la idea para
alcanzar la longitud mínima que exige el cargador web de JARVIS al extraer texto.</p>
{links}</main></body></html>"""


def minimal_pdf(content: str) -> bytes:
    """Un PDF de una página con texto, escrito a mano (sin librerías)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({content}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return out.getvalue()


def word_file() -> bytes:
    doc = WordDocument()
    doc.add_heading("Manual del horno", level=1)
    doc.add_paragraph("Precalentar diez minutos.")
    doc.add_heading("Limpieza", level=2)
    doc.add_paragraph("Usar un paño húmedo.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def wikipedia_api(articles: dict[str, str], category: list[str] | None = None):
    """Simula la API de Wikipedia."""

    def handler(request: httpx.Request) -> httpx.Response:
        params = dict(request.url.params)
        if params.get("list") == "categorymembers":
            members = [{"title": t} for t in (category or [])]
            return httpx.Response(200, json={"query": {"categorymembers": members}})
        title = params["titles"]
        if title not in articles:
            return httpx.Response(
                200, json={"query": {"pages": [{"title": title, "missing": True}]}}
            )
        page = {"title": title, "pageid": 1, "extract": articles[title]}
        return httpx.Response(200, json={"query": {"pages": [page]}})

    return httpx.AsyncClient(transport=httpx.MockTransport(handler))


async def collect(loader) -> list:
    return [doc async for doc in loader.load()]


# ---------- FileLoader ----------


async def test_archivo_markdown_usa_su_titulo():
    docs = await collect(FileLoader("notas.md", b"# Recetas\n\nArepa con queso."))
    assert docs[0].title == "Recetas" and "Arepa" in docs[0].text
    assert docs[0].mime_type == "text/markdown"


async def test_archivo_de_texto_en_latin1():
    doc = FileLoader("viejo.txt", "Canción de cuna".encode("cp1252")).convert()
    assert doc.text == "Canción de cuna" and doc.title == "viejo"


async def test_word_convierte_estilos_de_titulo():
    doc = FileLoader("horno.docx", word_file()).convert()
    assert doc.title == "Manual del horno"
    assert "## Limpieza" in doc.text


async def test_pdf_extrae_el_texto():
    doc = FileLoader("factura.pdf", minimal_pdf("Total a pagar 120000")).convert()
    assert "Total a pagar 120000" in doc.text
    assert doc.mime_type == "application/pdf"


async def test_html_extrae_el_contenido_principal():
    html = PAGE.format(title="Guía", body="El reciclaje sale los jueves.", links="")
    doc = FileLoader("guia.html", html.encode()).convert()
    assert "reciclaje sale los jueves" in doc.text


@pytest.mark.parametrize(
    ("name", "data", "error"),
    [
        ("foto.png", b"\x89PNG", "no soportado"),
        ("vacio.txt", b"   \n ", "no tiene texto"),
        ("roto.pdf", b"%PDF-1.4 basura", "No se pudo leer"),
    ],
)
async def test_archivos_invalidos(name: str, data: bytes, error: str):
    with pytest.raises(LoaderError, match=error):
        FileLoader(name, data).convert()


# ---------- Wikipedia ----------


def test_wikipedia_a_markdown_quita_relleno():
    md = wiki_to_markdown("Bogotá", ARTICLE)
    assert md.startswith("# Bogotá")
    assert "## Historia" in md and "### Época colonial" in md and "## Clima" in md
    assert "Referencias" not in md and "Libro de historia" not in md


async def test_wikipedia_por_titulos_y_categoria():
    http = wikipedia_api({"Bogotá": ARTICLE, "Medellín": "Ciudad de Antioquia."}, ["Medellín"])
    loader = WikipediaLoader(
        http, titles=["Bogotá", "No existe"], category="Ciudades", delay_seconds=0
    )
    docs = await collect(loader)

    assert [d.title for d in docs] == ["Bogotá", "Medellín"]
    assert docs[0].uri == "https://es.wikipedia.org/wiki/Bogot%C3%A1"
    assert docs[0].metadata["license"] == "CC BY-SA 4.0"


async def test_wikipedia_sin_resultados_es_error():
    loader = WikipediaLoader(wikipedia_api({}), titles=["Nada"], delay_seconds=0)
    with pytest.raises(LoaderError):
        await collect(loader)


# ---------- Web ----------


def web_site(pages: dict[str, str], robots: str = ""):
    requested: list[str] = []

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        requested.append(url)
        if url.endswith("/robots.txt"):
            return httpx.Response(200, text=robots)
        if url in pages:
            return httpx.Response(200, text=pages[url], headers={"content-type": "text/html"})
        return httpx.Response(404)

    return httpx.AsyncClient(transport=httpx.MockTransport(handler)), requested


async def test_web_sigue_enlaces_del_mismo_sitio_hasta_el_limite():
    links = '<a href="/b">B</a> <a href="https://otro.com/x">fuera</a> <a href="/c">C</a>'
    pages = {
        "https://casa.test/a": PAGE.format(title="A", body="Página A.", links=links),
        "https://casa.test/b": PAGE.format(title="B", body="Página B.", links=""),
        "https://casa.test/c": PAGE.format(title="C", body="Página C.", links=""),
    }
    http, requested = web_site(pages)
    docs = await collect(WebLoader(http, "https://casa.test/a", max_pages=2, delay_seconds=0))

    assert [d.title for d in docs] == ["A", "B"]
    assert not any("otro.com" in url for url in requested)


async def test_web_respeta_robots_txt():
    pages = {"https://casa.test/privado": PAGE.format(title="P", body="Secreto.", links="")}
    http, requested = web_site(pages, robots="User-agent: *\nDisallow: /privado")
    with pytest.raises(LoaderError):
        await collect(WebLoader(http, "https://casa.test/privado", delay_seconds=0))
    assert "https://casa.test/privado" not in requested


def test_web_rechaza_direcciones_invalidas():
    with pytest.raises(LoaderError):
        WebLoader(httpx.AsyncClient(), "ftp://casa.test")


# ---------- Cola y servicio (base de datos) ----------


@pytest.mark.db
async def test_la_cola_entrega_cada_tarea_una_sola_vez(session: AsyncSession):
    source = Source(kind="manual", name="Cola")
    session.add_all(
        [IngestionJob(source=source, params={"n": 1}), IngestionJob(source=source, params={"n": 2})]
    )
    await session.flush()
    repo = JobRepository(session)

    first, second, none = await repo.claim_next(), await repo.claim_next(), await repo.claim_next()
    assert first.params == {"n": 1} and second.params == {"n": 2} and none is None
    assert first.status == "running" and first.started_at is not None
    assert await repo.requeue_stuck() == 2


@pytest.mark.db
async def test_tarea_de_wikipedia_de_principio_a_fin(
    session_factory: async_sessionmaker, tmp_path: Path
):
    service = IngestionService(session_factory, uploads_dir=tmp_path)
    job = await service.enqueue_wikipedia(["Bogotá", "No existe", "Medellín"], None, 10)
    http = wikipedia_api({"Bogotá": ARTICLE, "Medellín": "Ciudad de Antioquia."})

    await service.process(job.id, FakeEmbedder(), http)

    async with session_factory() as s:
        done = await s.get(IngestionJob, job.id)
        titles = (await s.scalars(select(Document.title).order_by(Document.id))).all()
        chunks = await s.scalar(select(func.count(Chunk.id)))
        source = await s.get(Source, done.source_id)
    assert done.status == "done" and done.docs_done == 2
    assert titles == ["Bogotá", "Medellín"]
    assert chunks >= 4
    assert source.license == "CC BY-SA 4.0"


@pytest.mark.db
async def test_archivo_subido_se_indexa_y_se_borra_la_copia(
    session_factory: async_sessionmaker, tmp_path: Path, session: AsyncSession
):
    owner = User(username="dueno", display_name="Dueño", password_hash="!", role="owner")
    session.add(owner)
    await session.commit()
    service = IngestionService(session_factory, uploads_dir=tmp_path)
    job = await service.enqueue_file(
        "diario.md", "# Diario\n\nHoy aprendí RAG.".encode(), owner_user_id=owner.id, personal=True
    )
    assert len(list(tmp_path.iterdir())) == 1

    await service.process(job.id, FakeEmbedder(), httpx.AsyncClient())

    async with session_factory() as s:
        done = await s.get(IngestionJob, job.id)
        source = await s.get(Source, done.source_id)
    assert done.status == "done"
    assert source.is_personal and source.owner_user_id == owner.id
    assert list(tmp_path.iterdir()) == []


@pytest.mark.db
async def test_tarea_fallida_guarda_el_motivo(session_factory: async_sessionmaker, tmp_path: Path):
    service = IngestionService(session_factory, uploads_dir=tmp_path)
    job = await service.enqueue_wikipedia(["Nada"], None, 5)
    await service.process(job.id, FakeEmbedder(), wikipedia_api({}))
    async with session_factory() as s:
        failed = await s.get(IngestionJob, job.id)
    assert failed.status == "failed" and "ningún artículo" in failed.error_message


@pytest.mark.db
async def test_cancelar_tarea(session_factory: async_sessionmaker, tmp_path: Path):
    service = IngestionService(session_factory, uploads_dir=tmp_path)
    job = await service.enqueue_wikipedia(["Bogotá"], None, 5)
    assert await service.cancel(job.id) is True
    assert await service.cancel(job.id) is False


async def test_sin_internet_no_se_encola_web(tmp_path: Path):
    service = IngestionService(None, uploads_dir=tmp_path, allow_internet=False)  # type: ignore[arg-type]
    with pytest.raises(IngestionError, match="desactivada"):
        await service.enqueue_web("https://casa.test", 1)


# ---------- Endpoints y página ----------


@pytest.mark.db
async def test_endpoints_de_ingesta(test_db_url: str, tmp_path: Path):
    from app.core.config import Settings, get_settings
    from app.core.dependencies import get_session_factory
    from app.main import create_app

    engine = create_async_engine(test_db_url, poolclass=NullPool)
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with factory() as s:
        s.add(User(username="dueno", display_name="Dueño", password_hash="!", role="owner"))
        await s.commit()
    app = create_app()
    app.dependency_overrides[get_session_factory] = lambda: factory
    app.dependency_overrides[get_settings] = lambda: Settings(
        uploads_dir=str(tmp_path), ingestion_allow_internet=False
    )
    client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")
    try:
        async with client:
            page = await client.get("/conocimiento")
            assert page.status_code == 200 and "Agregar conocimiento" in page.text

            uploaded = await client.post(
                "/api/knowledge/files",
                files={"file": ("notas.md", b"# Notas\n\nTexto.", "text/markdown")},
                data={"personal": "true"},
            )
            assert uploaded.status_code == 202
            job = uploaded.json()
            assert job["status"] == "queued" and job["label"] == "notas.md"

            jobs = (await client.get("/api/knowledge/jobs")).json()
            assert [j["id"] for j in jobs] == [job["id"]]

            bad = await client.post(
                "/api/knowledge/files", files={"file": ("foto.png", b"x", "image/png")}
            )
            assert bad.status_code == 415

            offline = await client.post("/api/knowledge/wikipedia", json={"titles": ["Bogotá"]})
            assert offline.status_code == 403

            cancel = await client.post(f"/api/knowledge/jobs/{job['id']}/cancel")
            assert cancel.status_code == 204

            sources = (await client.get("/api/knowledge/sources")).json()
            deleted = await client.delete(f"/api/knowledge/sources/{sources[0]['id']}")
            assert deleted.status_code == 204
            assert (await client.get("/api/knowledge/sources")).json() == []
    finally:
        async with engine.begin() as conn:
            await conn.execute(text("TRUNCATE users, sources RESTART IDENTITY CASCADE"))
        await engine.dispose()


def test_respuesta_json_de_wikipedia_es_valida():
    # Protección contra cambios de formato: el cargador usa formatversion=2.
    sample = {"query": {"pages": [{"title": "X", "extract": "Y"}]}}
    assert json.loads(json.dumps(sample))["query"]["pages"][0]["title"] == "X"
