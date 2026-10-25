"""Carga archivos locales: texto, Markdown, PDF, Word y HTML."""

import io
from collections.abc import AsyncIterator
from pathlib import Path

import trafilatura
from docx import Document as WordDocument
from pypdf import PdfReader

from app.services.ingestion.base import LoaderError, RawDocument

MAX_BYTES = 20 * 1024 * 1024  # 20 MB por archivo
SUPPORTED = {".txt", ".md", ".markdown", ".pdf", ".docx", ".html", ".htm"}


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise LoaderError("No se pudo leer el texto del archivo")


def _pdf_to_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p.strip() for p in pages if p.strip())


def _docx_to_markdown(data: bytes) -> str:
    """Convierte los estilos "Título 1/2/3" de Word en títulos de Markdown."""
    lines = []
    for paragraph in WordDocument(io.BytesIO(data)).paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        level = next((n for n in (1, 2, 3) if style in (f"heading {n}", f"título {n}")), None)
        lines.append(f"{'#' * level} {text}" if level else text)
    return "\n\n".join(lines)


def _html_to_markdown(data: bytes) -> str:
    text = trafilatura.extract(
        _decode(data), output_format="markdown", include_tables=True, include_comments=False
    )
    return text or ""


def _first_heading(text: str) -> str | None:
    for line in text.splitlines():
        if line.startswith("# "):
            return line[2:].strip()
    return None


class FileLoader:
    """Lee un archivo (desde el disco o subido por la web) y lo convierte a Markdown."""

    def __init__(self, filename: str, data: bytes) -> None:
        self._filename = filename
        self._data = data

    @classmethod
    def from_path(cls, path: Path) -> "FileLoader":
        return cls(path.name, path.read_bytes())

    async def load(self) -> AsyncIterator[RawDocument]:
        yield self.convert()

    def convert(self) -> RawDocument:
        suffix = Path(self._filename).suffix.lower()
        if suffix not in SUPPORTED:
            raise LoaderError(f"Formato no soportado: {suffix or 'sin extensión'}")
        if len(self._data) > MAX_BYTES:
            raise LoaderError("El archivo supera 20 MB")
        try:
            if suffix == ".pdf":
                text, mime = _pdf_to_text(self._data), "application/pdf"
            elif suffix == ".docx":
                text = _docx_to_markdown(self._data)
                mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            elif suffix in (".html", ".htm"):
                text, mime = _html_to_markdown(self._data), "text/html"
            else:
                text = _decode(self._data)
                mime = "text/markdown" if suffix in (".md", ".markdown") else "text/plain"
        except LoaderError:
            raise
        except Exception as exc:  # archivos dañados de muchas formas distintas
            raise LoaderError(f"No se pudo leer {self._filename}: {exc}") from exc
        if not text.strip():
            raise LoaderError(f"{self._filename} no tiene texto (¿es un PDF escaneado?)")
        title = _first_heading(text) or Path(self._filename).stem.replace("_", " ")
        return RawDocument(title=title, text=text, uri=f"file:{self._filename}", mime_type=mime)
